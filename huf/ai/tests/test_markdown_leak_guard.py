"""Tests for the raw-markdown leak guard in the document render pipeline.

Run with:
	bench --site <site> run-tests --app huf --module huf.ai.tests.test_markdown_leak_guard
"""

import re
import time
import unittest
from io import BytesIO

from huf.ai.artifacts.render.html import (
	PRINT_STYLESHEET,
	render_document_html,
	render_document_html_with_report,
)
from huf.ai.artifacts.render.markdown_guard import (
	enable_markdown_in_containers,
	_fallback_text,
	fallback_markdown_leaks,
	find_markdown_leaks,
	repair_markdown_leaks,
)

SCREENSHOT = (
	'<div class="split"><div class="split-main">\n\n## Key Quarterly Achievements\n'
	"* **Enterprise Expansion:** Secured 14 new contracts.\n* **Core Platform:** done.\n\n"
	"| Region | Accounts | Status |\n| :--- | :---: | :---: |\n"
	'| **North America** | 84,200 | <span class="status-badge">EXCEEDED</span> |\n\n'
	'</div><div class="split-side">\n\n### Q1 2025 Priorities\n1. **Scale APAC Sales**\n2. Hire\n'
	"- [x] **SOC 2**\n\n</div></div>"
)

CODE_DOC = "<pre><code>## not a heading\n**not bold**\n| a | b |\n|---|---|</code></pre>"

CORPUS = {
	"nested_divs": '<div class="a"><div class="b">\n\n## Deep\n- one\n- two\n\n</div></div>',
	"table_in_split_main": '<div class="split"><div class="split-main">\n\n| A | B |\n|---|---|\n| **x** | 2 |\n\n</div></div>',
	"task_list": "<div>\n\n- [x] done **thing**\n- [ ] todo item\n\n</div>",
	"ordered": "<section>\n1. First\n2. Second\n3. Third\n</section>",
	"code_fence_with_md": "<div>\n\nText\n\n```\n## literal\n**literal**\n```\n\n- real item\n- another\n\n</div>",
	"inline_code_bold": "<div>\n\nUse `**kwargs` here and **bold** there.\n\n</div>",
	"blockquote_list": "<div>\n\n> quote\n\n- a\n- b\n\n</div>",
	"details": "<details><summary>More</summary>\n\n## Inside\n- x\n- y\n\n</details>",
	"callout_in_split": '<div class="split"><div class="split-side"><div class="callout">\n\n**Note:** read [the docs](https://example.com/a_b_c).\n\n</div></div></div>',
	"unclosed_div": '<div class="x">\n\n## Heading\n\n- a\n- b\n',
	"md_in_span": "<div><span>**bold in span**</span></div>",
	"heading_after_tag": "<div>\n## Direct heading\ntext\n</div>",
	"mixed": "<div>\nPlain <b>html</b> paragraph.\n\n## Then markdown\n\nmore **bold** text\n</div>",
	"arabic": '<div class="split-main">\n\n## النتائج\n- **المبيعات:** مرتفعة\n- الأرباح\n\n</div>',
	"emoji": "<div>\n\n## Wins \U0001F680\n- ✅ **Shipped**\n- \U0001F389 Party\n\n</div>",
	"long_table": "<div>\n\n| n | v |\n|---|---|\n" + "".join(f"| {i} | **{i*2}** |\n" for i in range(400)) + "\n</div>",
	"indented": "<div>\n    <div>\n\n    ## Indented\n    - a\n    - b\n\n    </div>\n</div>",
	"screenshot": SCREENSHOT,
}


def _body(html):
	return html[html.index("<body") :]


class TestRootFix(unittest.TestCase):
	def test_screenshot_renders_real_elements(self):
		body = _body(render_document_html(SCREENSHOT, language="html"))
		for needle in ("<h2>Key Quarterly", "<h3>Q1 2025", "<table>", "<ul>", "<ol>", "<strong>North America</strong>", '<span class="status-badge">EXCEEDED</span>'):
			self.assertIn(needle, body)
		self.assertEqual(find_markdown_leaks(body), [])
		self.assertNotIn("##", body)
		self.assertNotIn("**", body)

	def test_corpus_no_leaks(self):
		for name, src in CORPUS.items():
			for lang in ("html", "markdown"):
				with self.subTest(name=name, lang=lang):
					html, report = render_document_html_with_report(src, language=lang)
					self.assertEqual(find_markdown_leaks(_body(html)), [], name)
					self.assertEqual(report["leaks_remaining"], 0)

	def test_code_untouched(self):
		html = _body(render_document_html(CODE_DOC, language="html"))
		self.assertIn("## not a heading\n**not bold**\n| a | b |\n|---|---|", html)
		html = _body(render_document_html("<div>\n\n- a\n- b\n\nUse `**kwargs`\n\n</div>" + CODE_DOC, language="html"))
		self.assertIn("<code>**kwargs</code>", html)
		self.assertIn("## not a heading", html)
		md = _body(render_document_html("```\n## x\n**y**\n```\n\nUse `a **b** c`"))
		self.assertIn("## x\n**y**", md)
		self.assertIn("<code>a **b** c</code>", md)

	def test_attr_not_added_when_present_or_structural(self):
		src = '<div markdown="block">\n\n## H\n\n</div>'
		self.assertEqual(enable_markdown_in_containers(src), src)
		src = '<div class="doc-header">\n- a\n- b\n</div>'
		self.assertEqual(enable_markdown_in_containers(src), src)
		src = "<div>plain text only</div>"
		self.assertEqual(enable_markdown_in_containers(src), src)

	def test_idempotent(self):
		once = enable_markdown_in_containers(SCREENSHOT)
		self.assertNotEqual(once, SCREENSHOT)
		self.assertEqual(enable_markdown_in_containers(once), once)
		a = render_document_html(SCREENSHOT, language="html")
		self.assertEqual(a, render_document_html(SCREENSHOT, language="html"))

	def test_plain_html_unchanged(self):
		src = '<div class="doc-header"><span class="brand">Acme</span></div><p>5 * 3 = 15 in C# a_b_c</p><div class="metric-grid"><div class="metric">9</div></div>'
		self.assertEqual(enable_markdown_in_containers(src), src)
		body = _body(render_document_html(src, language="html"))
		self.assertIn(src, body)

	def test_text_not_mangled(self):
		src = "<div>\n\nCost 5 * 3 and 2 * 4, C# language, a_b_c, https://x.com/a_b_c_d and `**` in code\n\n- one\n- two\n\n</div>"
		body = _body(render_document_html(src, language="html"))
		self.assertIn("5 * 3 and 2 * 4", body)
		self.assertIn("C# language", body)
		self.assertIn("a_b_c", body)
		self.assertIn("https://x.com/a_b_c_d", body)

	def test_print_stylesheet_unchanged(self):
		self.assertIn("@page", PRINT_STYLESHEET)
		self.assertIn("running(foot)", PRINT_STYLESHEET)

	def test_performance(self):
		block = '<div class="split-main">\n\n## Heading\n- **a** item\n- b item\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n</div>\n'
		src = block * (200_000 // len(block))
		t = time.time()
		html = render_document_html(src, language="html")
		# Wall-clock guard against pathological (quadratic) blow-ups, not a benchmark: a 200 KB document
		# renders in ~1-3 s normally; the limit is generous so loaded CI runners do not flake.
		self.assertLess(time.time() - t, 20.0)
		self.assertEqual(find_markdown_leaks(_body(html)), [])


class TestDetector(unittest.TestCase):
	def test_false_positives(self):
		for text in (
			"<p>5 * 3 = 15</p>",
			"<p>I like C# and F#</p>",
			"<p>use snake_case and a_b_c here</p>",
			"<p>a | b only one pipe</p>",
			'<p>see <a href="https://x.com/a_b_c">https://x.com/a_b_c_d</a></p>',
			"<p>#hashtag and #1 trend</p>",
			"<p>2**3 and 4 ** 5</p>",
			"<pre><code>## h\n**b**\n- a\n- b</code></pre>",
			"<p>Use <code>**kwargs</code> and <code>## x</code></p>",
			"<p>- single dash line</p>",
			"<style>a > b { x: 1 } /* ## **x** */</style>",
		):
			with self.subTest(text=text):
				self.assertEqual(find_markdown_leaks(text), [])

	def test_true_positives(self):
		cases = {
			"heading": "<div>\n## Title\n</div>",
			"bullet-list": "<div>\n- a\n- b\n</div>",
			"ordered-list": "<div>\n1. a\n2. b\n</div>",
			"table": "<div>| a | b |\n|---|---|\n| 1 | 2 |</div>",
			"bold": "<p>some **bold** text</p>",
			"task-list": "<div>- [x] done</div>",
			"link": "<p>see [docs](https://example.com)</p>",
			"fence": "<div>```\ncode\n```</div>",
		}
		for kind, html in cases.items():
			with self.subTest(kind=kind):
				self.assertIn(kind, {lk.kind for lk in find_markdown_leaks(html)})


class TestRepair(unittest.TestCase):
	def test_repairs_leaked_html(self):
		leaked = '<div class="split-main">\n## Title\n- **a**\n- b\n\n| A | B |\n|---|---|\n| 1 | **2** |\n</div>'
		self.assertTrue(find_markdown_leaks(leaked))
		fixed, count = repair_markdown_leaks(leaked)
		self.assertGreaterEqual(count, 1)
		self.assertEqual(find_markdown_leaks(fixed), [])
		for needle in ("<h2>Title</h2>", "<strong>a</strong>", "<li>", "<table>", "<strong>2</strong>"):
			self.assertIn(needle, fixed)
		self.assertIn('class="split-main"', fixed)

	def test_repair_keeps_sanitizer(self):
		leaked = "<div>\n## T\n- a\n- b\n<script>alert(1)</script>\n</div>"
		fixed, _ = repair_markdown_leaks(leaked)
		self.assertNotIn("<script", fixed)

	def test_noop_on_clean(self):
		clean = _body(render_document_html(SCREENSHOT, language="html"))
		out, count = repair_markdown_leaks(clean)
		self.assertEqual(out, clean)
		self.assertEqual(count, 0)

	def test_report_shape(self):
		_, report = render_document_html_with_report(SCREENSHOT, language="html")
		self.assertEqual(report["leaks_remaining"], 0)
		self.assertIsInstance(report["kinds"], list)


class TestDocxNoLiteralMarkdown(unittest.TestCase):
	def test_docx_from_leaking_fixture(self):
		from docx import Document

		from huf.ai.artifacts.render.docx import html_to_docx

		html = render_document_html(SCREENSHOT, language="html")
		doc = Document(BytesIO(html_to_docx(html)))
		texts = [p.text for p in doc.paragraphs]
		for t in doc.tables:
			for row in t.rows:
				for cell in row.cells:
					texts.extend(p.text for p in cell.paragraphs)
		joined = "\n".join(texts)
		self.assertIn("Key Quarterly Achievements", joined)
		self.assertNotIn("##", joined)
		self.assertNotIn("**", joined)
		self.assertNotIn("| :---", joined)


class TestPromptGuard(unittest.TestCase):
	def test_rule_text_present(self):
		from huf.ai.document_artifact_instructions import DOCUMENT_ARTIFACT_INSTRUCTIONS as I

		self.assertIn('markdown="1"', I)
		self.assertIn("NEVER write raw markdown", I)
		self.assertIn(":::columns-2", I)
		self.assertIn("BLANK LINE before and after every heading, list, table and code fence", I)
		self.assertIn("Wrong (no blank lines", I)


def _plain(html):
	return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html.replace("<br>", " "))).strip()


class TestFallbackSpacing(unittest.TestCase):
	REPRO = (
		"<div><span>## Section Heading **bold bit** and more\n- first item\n- second item\n"
		"| Name | Score |\n| --- | --- |\n| Ann | **9** |\n</span></div>"
	)

	def test_repro_keeps_spaces_and_blocks(self):
		out, count = fallback_markdown_leaks(
			"<div><span>## Section Heading <strong>bold bit</strong> and more\n- first item\n- second item\n"
			"| Name | Score |\n| --- | --- |\n| Ann | 9 |\n</span></div>"
		)
		self.assertGreaterEqual(count, 1)
		self.assertIn("<strong>Section Heading</strong>", out)
		self.assertRegex(out, r"<strong>Section Heading</strong>\s*<br>")
		plain = _plain(out)
		self.assertIn("Section Heading bold bit and more", plain)
		self.assertIn("\u2022 first item", plain)
		self.assertIn("\u2022 second item", plain)
		self.assertIn("Name Score", plain)
		self.assertIn("Ann 9", plain)
		for bad in ("##", "**", "---"):
			self.assertNotIn(bad, out)

	def test_full_pipeline_repro(self):
		out = _body(render_document_html(self.REPRO, language="html"))
		plain = _plain(out)
		self.assertIn("Section Heading bold bit and more", plain)
		for bad in ("##", "**", "---"):
			self.assertNotIn(bad, plain)
		self.assertEqual(find_markdown_leaks(out), [])

	def test_inline_markers_keep_surrounding_spaces(self):
		t = _fallback_text("a **b** c __d__ e [t](https://x.io) f")
		self.assertEqual(t, "a b c d e t (https://x.io) f")

	def test_no_content_loss_and_idempotent(self):
		sources = [
			"## Title one\n- **alpha** beta\n- gamma [delta](https://e.io)\n",
			"1. first **x** thing\n2. second\n| A | B |\n|---|---|\n| one | **two** |\n",
			"- [x] done task\n- [ ] open task\n",
		]
		for src in sources:
			out = _fallback_text(src)
			cleaned = re.sub(r"\*\*|__|#+|\||[-*+]\s|\d[.)]\s|\[[ xX]\]|\]\(|[\[\]()]|:?-{3,}:?", " ", src)
			want = [w for w in cleaned.split() if re.search(r"\w", w)]
			got = _plain(out)
			pos = 0
			for w in want:
				i = got.find(w, pos)
				self.assertGreaterEqual(i, 0, (w, got))
				pos = i + len(w)
			for bad in ("##", "**", "|"):
				self.assertNotIn(bad, out)
			# Re-running the fallback over its own output is a no-op.
			wrapped = "<div>" + out + "</div>"
			self.assertEqual(fallback_markdown_leaks(wrapped), (wrapped, 0))
