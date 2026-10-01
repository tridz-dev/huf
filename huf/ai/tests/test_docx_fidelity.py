# Copyright (c) 2026, Tridz Technologies Pvt Ltd and contributors
# For license information, please see license.txt

"""DOCX fidelity: the .docx a user downloads looks like the Document view, and a
format preview is the exact file Export saves.

Renders one realistic sample (headings, inline formatting, nested and ordered
lists, a table with column alignment, a code block, a quote, a callout, a rule
followed by more content, a link) through the real pipeline and checks the
OOXML for the things that made the old output look poor: flattened inline
formatting, style-only bullets and borders, theme fonts, no page numbers, raw
markdown, and content silently dropped after a horizontal rule.

Run with:
	bench --site <site> run-tests --app huf --module huf.ai.tests.test_docx_fidelity
"""

import base64
import hashlib
import io
import re
import unittest
import zipfile

import frappe

from huf.ai.artifact_export_api import export_document_content, preview_document_file
from huf.ai.artifacts.render.docx import BODY_FONT, HEADING_FONT, MONO_FONT, html_to_docx
from huf.ai.artifacts.render.html import render_document_html

SAMPLE = """# Q3 Operations Review

A short report with **bold**, *italic*, `inline code` and a [link](https://example.com/q3).

## Highlights

- Revenue grew **18%** quarter over quarter
- Churn fell to *2.1%*
    - Enterprise churn: 0.8%
    - SMB churn: 3.4%
- Support backlog cleared

## Next steps

1. Ship the billing migration
2. Hire two support engineers
3. Review pricing in October

## Numbers

| Region | Revenue | Growth |
|--------|--------:|-------:|
| North  | $1.2M   | 12%    |
| South  | $0.9M   | 22%    |

## Code

```python
def growth(prev, cur):
    return (cur - prev) / prev
```

> Quality is not an act, it is a habit.

<div class="callout"><strong>Summary:</strong> the quarter closed ahead of plan.</div>

---

Final paragraph after the rule.
"""


def _parts(data: bytes) -> dict[str, str]:
	with zipfile.ZipFile(io.BytesIO(data)) as z:
		return {n: z.read(n).decode("utf-8") for n in z.namelist() if n.endswith((".xml", ".rels"))}


def _paragraph_text(p_xml: str) -> str:
	return "".join(re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>", p_xml))


class TestDocxFidelity(unittest.TestCase):
	@classmethod
	def setUpClass(cls):
		cls.html = render_document_html(SAMPLE, title="Q3 Operations Review")
		cls.data = html_to_docx(cls.html)
		cls.parts = _parts(cls.data)
		cls.body = cls.parts["word/document.xml"]
		cls.paragraphs = re.findall(r"<w:p[ >].*?</w:p>", cls.body, re.S)

	def test_opens_with_python_docx(self):
		from docx import Document

		doc = Document(io.BytesIO(self.data))
		self.assertEqual(doc.core_properties.title, "Q3 Operations Review")

	def test_no_content_lost_after_horizontal_rule(self):
		# `---` used to open an <hr> node that swallowed the rest of the document.
		self.assertIn("Final paragraph after the rule.", self.body)
		self.assertNotIn("________", self.body)
		self.assertIn('<w:bottom w:val="single"', self.body)

	def test_no_raw_markdown_leaks(self):
		text = " ".join(_paragraph_text(p) for p in self.paragraphs)
		for marker in ("**", "```", "](", "|---", "- Enterprise", "1. Ship"):
			self.assertNotIn(marker, text, marker)

	def test_inline_formatting_is_real_runs(self):
		intro = next(p for p in self.paragraphs if "A short report" in p)
		self.assertIn("<w:b/>", intro)
		self.assertIn("<w:i/>", intro)
		self.assertIn(f'w:ascii="{MONO_FONT}"', intro)
		self.assertIn("<w:hyperlink", intro)
		rels = self.parts["word/_rels/document.xml.rels"]
		self.assertIn('Target="https://example.com/q3"', rels)
		self.assertIn('TargetMode="External"', rels)

	def test_lists_use_real_numbering_with_nesting(self):
		numbered = [p for p in self.paragraphs if "<w:numPr>" in p]
		self.assertEqual(len(numbered), 8)
		nested = [p for p in numbered if '<w:ilvl w:val="1"/>' in p]
		self.assertEqual([_paragraph_text(p) for p in nested], ["Enterprise churn: 0.8%", "SMB churn: 3.4%"])
		self.assertEqual(_paragraph_text(numbered[1]), "Churn fell to 2.1%")
		numbering = self.parts["word/numbering.xml"]
		self.assertIn('w:lvlText w:val="•"', numbering)
		self.assertIn('<w:numFmt w:val="decimal"/>', numbering)
		self.assertIn("<w:startOverride", numbering)

	def test_table_has_explicit_borders_header_and_alignment(self):
		table = re.search(r"<w:tbl>.*?</w:tbl>", self.body, re.S).group(0)
		self.assertIn("<w:tblBorders>", table)
		self.assertIn('<w:insideH w:val="single"', table)
		self.assertIn("<w:tblHeader/>", table)
		self.assertIn('w:fill="F7FAFD"', table)  # --surface header fill
		header = re.search(r"<w:tr[ >].*?</w:tr>", table, re.S).group(0)
		self.assertIn("<w:b/>", header)
		self.assertIn('<w:jc w:val="right"/>', table)  # the Revenue/Growth columns

	def test_code_block_keeps_lines_and_mono_font(self):
		code = next(p for p in self.paragraphs if "def growth" in p)
		self.assertIn(f'w:ascii="{MONO_FONT}"', code)
		self.assertIn("<w:br/>", code)
		self.assertIn("    return (cur - prev) / prev", code)
		self.assertIn("<w:shd", code)

	def test_callout_keeps_bold_label(self):
		callout = re.findall(r"<w:tbl>.*?</w:tbl>", self.body, re.S)[-1]
		self.assertIn("Summary:", callout)
		self.assertIn("<w:b/>", callout)

	def test_styles_use_explicit_fonts_not_theme(self):
		styles = self.parts["word/styles.xml"]
		for style_id, font in (("Normal", BODY_FONT), ("Heading1", HEADING_FONT), ("Heading2", HEADING_FONT)):
			style = re.search(rf'<w:style [^>]*w:styleId="{style_id}".*?</w:style>', styles, re.S).group(0)
			self.assertIn(f'w:ascii="{font}"', style, style_id)
			self.assertNotIn("asciiTheme", style, style_id)
		heading1 = re.search(r'<w:style [^>]*w:styleId="Heading1".*?</w:style>', styles, re.S).group(0)
		self.assertIn('<w:sz w:val="56"/>', heading1)  # 28pt, like the HTML h1

	def test_page_is_a4_with_2cm_margins_and_page_numbers(self):
		sect = re.search(r"<w:sectPr.*?</w:sectPr>", self.body, re.S).group(0)
		self.assertIn('w:w="11906"', sect)
		self.assertIn('w:h="16838"', sect)
		self.assertIn('w:left="1134"', sect)
		footer = "".join(v for k, v in self.parts.items() if k.startswith("word/footer"))
		self.assertIn(" PAGE ", footer)
		self.assertIn(" NUMPAGES ", footer)

	def test_same_input_same_bytes(self):
		self.assertEqual(html_to_docx(self.html), self.data)


class TestFormatPreviewIsTheExportFile(unittest.TestCase):
	"""The preview for a format must be rendered from the same bytes Export saves."""

	def setUp(self):
		frappe.set_user("Administrator")

	def _pair(self, fmt: str) -> tuple[dict, dict]:
		preview = preview_document_file(SAMPLE, fmt, "markdown", "Q3 Operations Review")
		export = export_document_content(SAMPLE, fmt, "markdown", "Q3 Operations Review")
		return preview, export

	def test_docx_preview_bytes_hash_equal_export_bytes(self):
		preview, export = self._pair("docx")
		preview_bytes = base64.b64decode(preview["content_base64"])
		export_bytes = base64.b64decode(export["content_base64"])
		self.assertEqual(hashlib.sha256(preview_bytes).hexdigest(), hashlib.sha256(export_bytes).hexdigest())
		self.assertEqual(preview["sha256"], hashlib.sha256(preview_bytes).hexdigest())
		self.assertEqual(preview["sha256"], export["sha256"])
		self.assertEqual(preview["file_name"], export["file_name"])
		self.assertEqual(preview["file_name"], "Q3 Operations Review.docx")
		self.assertEqual(preview["mime_type"], export["mime_type"])
		self.assertTrue(preview_bytes.startswith(b"PK"))

	def test_html_and_md_preview_equal_export(self):
		for fmt in ("html", "md"):
			preview, export = self._pair(fmt)
			self.assertEqual(preview["content_base64"], export["content_base64"], fmt)

	def test_pdf_preview_bytes_hash_equal_export_bytes(self):
		# WeasyPrint writes no creation date or random file ID for these
		# documents (measured on the bench, 2026-09-30), so PDF holds too.
		preview, export = self._pair("pdf")
		data = base64.b64decode(preview["content_base64"])
		self.assertTrue(data.startswith(b"%PDF"))
		self.assertEqual(preview["sha256"], hashlib.sha256(data).hexdigest())
		self.assertEqual(preview["sha256"], export["sha256"])

	def test_rejects_unknown_format_and_empty_content(self):
		with self.assertRaises(frappe.ValidationError):
			preview_document_file(SAMPLE, "zip")
		with self.assertRaises(frappe.ValidationError):
			preview_document_file("", "docx")
